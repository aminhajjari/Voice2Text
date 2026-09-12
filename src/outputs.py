import os
import uuid
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
except ImportError:
    Document = None

from .utils import format_timestamp


def _atomic_write_text(output_path: Path, content: str, encoding: str = "utf-8"):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.parent / f".{output_path.name}.tmp.{uuid.uuid4().hex}"
    try:
        temp_path.write_text(content, encoding=encoding)
        temp_path.replace(output_path)
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def _atomic_save_docx(doc, output_path: Path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.parent / f".{output_path.name}.tmp.{uuid.uuid4().hex}"
    try:
        doc.save(str(temp_path))
        temp_path.replace(output_path)
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def _extract_transcript_data(transcription, default_filename="transcript"):
    """
    Extract (text, segments, filename) from either a dict or a MeetingDocument.
    Preserves exact backward compatibility for dict and str inputs.
    """
    if hasattr(transcription, "transcript_segments"):
        # MeetingDocument instance
        doc = transcription
        segments = [
            {"start": seg.start_time, "end": seg.end_time, "text": seg.text}
            for seg in doc.transcript_segments
        ]
        text = ""
        if hasattr(doc, "summary") and getattr(doc.summary, "text", None):
            text = doc.summary.text
        elif segments:
            text = "\n".join(seg["text"] for seg in segments if seg.get("text")).strip()

        filename = default_filename
        if hasattr(doc, "source_audio") and getattr(doc.source_audio, "original_name", None):
            filename = doc.source_audio.original_name

        return text, segments, filename

    if isinstance(transcription, dict):
        text = transcription.get("text", "")
        segments = transcription.get("segments", [])
        filename = transcription.get("filename", default_filename)
        return text, segments, filename

    text = str(transcription)
    return text, [], default_filename


def save_txt(transcription, output_path):
    output_path = Path(output_path)
    content, _, _ = _extract_transcript_data(transcription, default_filename=output_path.name)
    _atomic_write_text(output_path, content, encoding="utf-8")


def save_docx(transcription, output_path):
    if Document is None:
        raise RuntimeError("python-docx is required. Install it with: pip install python-docx")

    output_path = Path(output_path)
    content, _, filename = _extract_transcript_data(transcription, default_filename=output_path.name)

    doc = Document()
    doc.add_heading(output_path.stem.replace("_", " ").replace("-", " "), level=1)
    doc.add_paragraph(f"Filename: {filename}")
    doc.add_paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("Transcription:")
    for line in content.splitlines():
        if line.strip():
            doc.add_paragraph(line)

    _atomic_save_docx(doc, output_path)


def save_srt(transcription, output_path):
    output_path = Path(output_path)
    _, segments, _ = _extract_transcript_data(transcription, default_filename=output_path.name)

    lines = []
    subtitle_idx = 1
    for segment in segments:
        start = format_timestamp(segment.get("start", 0))
        end = format_timestamp(segment.get("end", 0))
        text = (segment.get("text", "") or "").strip().replace("\n", " ")
        if text:
            lines.append(str(subtitle_idx))
            lines.append(f"{start} --> {end}")
            lines.append(text)
            lines.append("")
            subtitle_idx += 1

    content = ("\n".join(lines).rstrip() + "\n") if lines else ""
    _atomic_write_text(output_path, content, encoding="utf-8")

def save_json(transcription, output_path, audio_path=None):
    from .models.meeting import MeetingDocument

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    if isinstance(transcription, MeetingDocument):
        doc = transcription
    elif isinstance(transcription, dict):
        file_path_str = str(audio_path) if audio_path else transcription.get("filename", str(output_path.stem))
        doc = MeetingDocument.from_transcription_result(transcription, file_path=file_path_str)
    else:
        file_path_str = str(audio_path) if audio_path else str(output_path.stem)
        doc = MeetingDocument.from_transcription_result({"text": str(transcription), "segments": []}, file_path=file_path_str)

    _atomic_write_text(output_path, doc.to_json(), encoding="utf-8")
