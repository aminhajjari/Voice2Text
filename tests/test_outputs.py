import tempfile
import unittest
from pathlib import Path

from src.outputs import save_srt


class TestOutputs(unittest.TestCase):
    def test_save_srt_sequential_indices_when_segments_empty(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "sub.srt"
            transcription = {
                "segments": [
                    {"start": 0.0, "end": 1.5, "text": "First line"},
                    {"start": 1.5, "end": 2.0, "text": ""},  # empty segment
                    {"start": 2.0, "end": 3.5, "text": "Second line"},
                ]
            }
            save_srt(transcription, output_path)
            content = output_path.read_text(encoding="utf-8")

            # Check that the second valid block gets index '2' instead of '3'
            self.assertIn("1\n00:00:00,000 --> 00:00:01,500\nFirst line", content)
            self.assertIn("2\n00:00:02,000 --> 00:00:03,500\nSecond line", content)
            self.assertNotIn("\n3\n", content)


    def test_save_json_meeting_document(self):
        import json
        from src.outputs import save_json
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "sub.json"
            transcription = {
                "text": "سلام دنیا",
                "segments": [
                    {"start": 0.0, "end": 1.5, "text": "سلام"},
                    {"start": 1.5, "end": 3.0, "text": "دنیا"},
                ],
                "filename": "audio.wav",
                "duration": 3.0,
            }
            save_json(transcription, output_path)
            content = output_path.read_text(encoding="utf-8")
            data = json.loads(content)
            self.assertEqual(data["summary"]["text"], "سلام دنیا")
            self.assertEqual(len(data["transcript_segments"]), 2)
            self.assertEqual(data["transcript_segments"][0]["text"], "سلام")

    def test_atomic_write_creates_new_file_and_replaces_existing(self):
        from src.outputs import save_txt
        with tempfile.TemporaryDirectory() as temp_dir:
            out_file = Path(temp_dir) / "test.txt"

            # 1. Successful creation
            save_txt({"text": "Version 1"}, out_file)
            self.assertTrue(out_file.exists())
            self.assertEqual(out_file.read_text(encoding="utf-8"), "Version 1")

            # 2. Successful atomic replacement of existing file
            save_txt({"text": "Version 2"}, out_file)
            self.assertEqual(out_file.read_text(encoding="utf-8"), "Version 2")

            # Verify no temporary files remain in directory
            temp_files = list(Path(temp_dir).glob(".*.tmp.*"))
            self.assertEqual(len(temp_files), 0)

    def test_atomic_write_cleans_up_on_failure_and_leaves_original_intact(self):
        from unittest.mock import patch
        from src.outputs import save_txt
        with tempfile.TemporaryDirectory() as temp_dir:
            out_file = Path(temp_dir) / "original.txt"
            out_file.write_text("Original pristine content", encoding="utf-8")

            # Simulate failure during Path.replace
            with patch("pathlib.Path.replace", side_effect=PermissionError("Simulated locked file")):
                with self.assertRaises(PermissionError):
                    save_txt({"text": "Corrupted attempt"}, out_file)

            # Original destination must remain unchanged
            self.assertEqual(out_file.read_text(encoding="utf-8"), "Original pristine content")

            # No orphaned temporary files left behind
            temp_files = list(Path(temp_dir).glob(".*.tmp.*"))
            self.assertEqual(len(temp_files), 0)

    def test_exporters_support_meeting_document(self):
        from src.models.meeting import MeetingDocument, TranscriptSegment, AudioInfo, Summary
        from src.outputs import save_txt, save_docx, save_srt

        meeting = MeetingDocument(
            id="meeting-test-doc",
            source_audio=AudioInfo(
                file_path="audio/test.wav",
                original_name="test.wav",
                size_bytes=1000,
                duration_seconds=5.0,
                format="wav",
            ),
            metadata={},
            transcript_segments=[
                TranscriptSegment(id="seg-1", start_time=0.0, end_time=2.0, text="سلام"),
                TranscriptSegment(id="seg-2", start_time=2.0, end_time=4.5, text="خوش آمدید"),
            ],
            summary=Summary(text="سلام\nخوش آمدید"),
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            txt_path = Path(temp_dir) / "out.txt"
            srt_path = Path(temp_dir) / "out.srt"
            docx_path = Path(temp_dir) / "out.docx"

            # 1. TXT export from MeetingDocument
            save_txt(meeting, txt_path)
            self.assertTrue(txt_path.exists())
            self.assertIn("سلام", txt_path.read_text(encoding="utf-8"))
            self.assertIn("خوش آمدید", txt_path.read_text(encoding="utf-8"))

            # 2. SRT export from MeetingDocument
            save_srt(meeting, srt_path)
            self.assertTrue(srt_path.exists())
            srt_content = srt_path.read_text(encoding="utf-8")
            self.assertIn("00:00:00,000 --> 00:00:02,000", srt_content)
            self.assertIn("سلام", srt_content)
            self.assertIn("00:00:02,000 --> 00:00:04,500", srt_content)
            self.assertIn("خوش آمدید", srt_content)

            # 3. DOCX export from MeetingDocument
            save_docx(meeting, docx_path)
            self.assertTrue(docx_path.exists())
            self.assertTrue(docx_path.stat().st_size > 0)
            from docx import Document
            doc = Document(str(docx_path))
            full_docx_text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("Filename: test.wav", full_docx_text)
            self.assertIn("سلام", full_docx_text)
            self.assertIn("خوش آمدید", full_docx_text)


if __name__ == "__main__":
    unittest.main()
